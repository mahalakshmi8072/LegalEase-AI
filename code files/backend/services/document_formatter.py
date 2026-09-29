import html
import re
from io import BytesIO
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF

ASSETS_DIR = Path(__file__).resolve().parents[2] / "assets"
FONT_DIR = ASSETS_DIR / "fonts"
LOGO_PATH = ASSETS_DIR / "logo.png"
DEJAVU_REGULAR = FONT_DIR / "DejaVuSans.ttf"
DEJAVU_BOLD = FONT_DIR / "DejaVuSans-Bold.ttf"

def sanitize_text(text: str) -> str:
    replacements = {"\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
                    "\u2013": "-", "\u2014": "-", "\u00a0": " ", "\u2022": "-"}
    for old, new in replacements.items():
        text = text.replace(old, new)
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text).strip()

def _split_lines(text: str):
    return [line.strip() for line in sanitize_text(text).splitlines()]

def format_html_preview(text: str) -> str:
    safe = html.escape(sanitize_text(text))
    safe = re.sub(r"\n{2,}", "</p><p>", safe)
    safe = safe.replace("\n", "<br>")
    return '<div class="legal-preview"><p>' + safe + "</p></div>"

def format_docx(text: str, doc_type: str, logo_path: Path = LOGO_PATH) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7); section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85); section.right_margin = Inches(0.85)
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)

        # LegalEase logo/header
    logo_para = doc.add_paragraph()
    logo_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    scale_run = logo_para.add_run("⚖ ")
    scale_run.font.name = "Segoe UI Symbol"
    scale_run.font.size = Pt(18)

    brand_run = logo_para.add_run("LegalEase")
    brand_run.bold = True
    brand_run.font.name = "Arial"
    brand_run.font.size = Pt(18)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.size = Pt(16)

   

    for line in _split_lines(text):
        if not line:
            doc.add_paragraph(); continue
         # Skip the title if Gemini already included it
        if line.strip().upper() == doc_type.strip().upper():continue
        heading = bool(re.match(r"^(\d+\.|SECTION\s+\d+|[A-Z][A-Z\s&-]{5,})$", line))
        p = doc.add_paragraph()
        r = p.add_run(line); r.bold = heading; r.font.name = "Times New Roman"
        r.font.size = Pt(12 if heading else 11)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase | AI-generated draft — review with a qualified lawyer").font.size = Pt(8)

    out = BytesIO(); doc.save(out); return out.getvalue()

class BrandedPDF(FPDF):
    def __init__(self, doc_type: str, logo_path: Path):
        super().__init__()
        self.doc_type = doc_type; self.logo_path = logo_path
        self.set_auto_page_break(auto=True, margin=18)
        if DEJAVU_REGULAR.exists() and DEJAVU_BOLD.exists():
            self.add_font("DejaVu", "", str(DEJAVU_REGULAR))
            self.add_font("DejaVu", "B", str(DEJAVU_BOLD))
            self.font_family = "DejaVu"
        else:
            self.font_family = "Helvetica"

    def header(self):
        if self.logo_path.exists():
            self.image(str(self.logo_path), x=92, y=8, w=26); self.set_y(38)
        else:
            self.set_y(12)
        self.set_font(self.font_family, "B", 12)
        self.cell(0, 7, self.doc_type.upper(), align="C"); self.ln(10)

    def footer(self):
        self.set_y(-14); self.set_font(self.font_family, "", 8)
        self.cell(0, 8, f"LegalEase | Page {self.page_no()} | AI-generated draft — legal review recommended", align="C")

def format_pdf(text: str, doc_type: str, logo_path: Path = LOGO_PATH) -> bytes:
    pdf = BrandedPDF(doc_type, logo_path); pdf.add_page()
    pdf.set_font(pdf.font_family, "", 11)
    for line in _split_lines(text):
        if not line: pdf.ln(4); continue
        heading = bool(re.match(r"^(\d+\.|SECTION\s+\d+|[A-Z][A-Z\s&-]{5,})$", line))
        pdf.set_font(pdf.font_family, "B" if heading else "", 12 if heading else 11)
        pdf.multi_cell(0, 6, line); pdf.ln(2)
    return bytes(pdf.output())
